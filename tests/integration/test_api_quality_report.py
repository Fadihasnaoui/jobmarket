"""Integration tests for POST /cv/{id}/quality-report.

Same DB-backed fixture convention as `test_api_improve.py`. All LLM calls mocked at
`jobmarket.api.quality_routes.generate_cv_quality_report`/`generate_improved_cv` —
zero network access.
"""

from __future__ import annotations

import base64

import pytest
from httpx import ASGITransport, AsyncClient
from test_api_cv import FakeEncoder, _add_job, _add_run, _fixture, _strong_cv_text

from jobmarket.api.app import create_app
from jobmarket.api.dependencies import get_cv_store, get_upload_encoder
from jobmarket.api.store import CvStore
from jobmarket.cv.extraction import extract_cv_profile
from jobmarket.cv.improver import CvImproverError, ImprovedCvResult, ImprovedExperienceEntry
from jobmarket.cv.llm_extraction import CvExtractionResult
from jobmarket.cv.quality_report import CvQualityReport, QualityCheck
from jobmarket.db.session import async_session_factory


def _fake_extract_deterministic(
    document: object,
    *,
    ontology: object = None,
    skill_ids: object = None,
    client: object = None,
):
    profile = extract_cv_profile(document, ontology=ontology, skill_ids=skill_ids)
    return CvExtractionResult(profile, [], "deterministic_fallback", "forced for test")


def _fake_improved_result() -> ImprovedCvResult:
    return ImprovedCvResult(
        improved_summary="Reworded summary using Python and AWS.",
        improved_experiences=[
            ImprovedExperienceEntry(
                title="Python Developer",
                employer="ExampleCo",
                start_date="Jan 2023",
                end_date="Dec 2023",
                entry_type="job",
                original_evidence="Built APIs with Python and deployed on AWS.",
                improved_description="Engineered APIs in Python and deployed them on AWS.",
            )
        ],
        skills_section=["Python", "AWS"],
        recommended_skills_to_develop=[],
        changes_explanation=[],
        source_degraded=False,
        source_degraded_reason=None,
    )


def _fake_report(*, contact_ok: bool) -> CvQualityReport:
    contact_check = QualityCheck(
        id="contact_info",
        label="Contact Information",
        status="pass" if contact_ok else "fail",
        score=10.0 if contact_ok else 0.0,
        max_score=10.0,
        message="ok" if contact_ok else "missing",
        suggestion=None if contact_ok else "Add your name and contact info.",
    )
    return CvQualityReport(
        overall_score=70 if contact_ok else 60,
        checks=[contact_check],
        source_degraded=False,
    )


@pytest.fixture
def app_client() -> tuple[object, CvStore]:
    app = create_app()
    store = CvStore()
    app.dependency_overrides[get_cv_store] = lambda: store
    app.dependency_overrides[get_upload_encoder] = lambda: FakeEncoder()
    return app, store


async def test_quality_report_returns_scored_result_and_docx(
    monkeypatch: pytest.MonkeyPatch,
    app_client: tuple[object, CvStore],
) -> None:
    app, _store = app_client
    monkeypatch.setattr(
        "jobmarket.api.cv_routes.extract_cv_profile_with_fallback", _fake_extract_deterministic
    )
    monkeypatch.setattr(
        "jobmarket.api.quality_routes.generate_cv_quality_report",
        lambda profile, skill_gap, matched_jobs, **kw: _fake_report(contact_ok=True),
    )
    monkeypatch.setattr(
        "jobmarket.api.quality_routes.generate_improved_cv",
        lambda profile, skill_gap, matched_jobs: _fake_improved_result(),
    )

    async with _fixture() as fixture, async_session_factory() as session:
        run_id = await _add_run(session, fixture)
        await _add_job(session, fixture, "Python Developer", ["Python", "AWS"], run_id)
        await session.commit()

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            upload_resp = await client.post(
                "/cv/upload",
                files={"file": ("cv.txt", _strong_cv_text().encode(), "text/plain")},
            )
            cv_id = upload_resp.json()["cv_id"]

            resp = await client.post(
                f"/cv/{cv_id}/quality-report",
                params={"run_id": run_id},
                json={"candidate_name": "Jane Doe", "contact_info": "jane@example.com"},
            )
            assert resp.status_code == 200
            payload = resp.json()

            assert payload["out_of_scope"] is False
            assert payload["overall_score"] == 70
            assert payload["checks"][0]["id"] == "contact_info"
            assert payload["checks"][0]["status"] == "pass"

            docx_bytes = base64.b64decode(payload["docx_base64"])
            assert docx_bytes[:2] == b"PK"  # docx is a zip container

            import io

            from docx import Document

            document = Document(io.BytesIO(docx_bytes))
            all_text = "\n".join(p.text for p in document.paragraphs)
            assert "Jane Doe" in all_text


async def test_quality_report_is_not_stale_across_calls_with_different_contact(
    monkeypatch: pytest.MonkeyPatch,
    app_client: tuple[object, CvStore],
) -> None:
    """The contact check must reflect each call's own contact_info, never a cached
    value from an earlier call — this is the whole reason the report isn't cached."""
    app, _store = app_client
    monkeypatch.setattr(
        "jobmarket.api.cv_routes.extract_cv_profile_with_fallback", _fake_extract_deterministic
    )

    def _report_from_contact(
        profile: object, skill_gap: object, matched_jobs: object, **kw: object
    ):
        return _fake_report(contact_ok=bool(kw.get("contact_info")))

    monkeypatch.setattr(
        "jobmarket.api.quality_routes.generate_cv_quality_report", _report_from_contact
    )
    monkeypatch.setattr(
        "jobmarket.api.quality_routes.generate_improved_cv",
        lambda profile, skill_gap, matched_jobs: _fake_improved_result(),
    )

    async with _fixture() as fixture, async_session_factory() as session:
        run_id = await _add_run(session, fixture)
        await _add_job(session, fixture, "Python Developer", ["Python", "AWS"], run_id)
        await session.commit()

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            upload_resp = await client.post(
                "/cv/upload",
                files={"file": ("cv.txt", _strong_cv_text().encode(), "text/plain")},
            )
            cv_id = upload_resp.json()["cv_id"]

            first = await client.post(f"/cv/{cv_id}/quality-report", params={"run_id": run_id})
            second = await client.post(
                f"/cv/{cv_id}/quality-report",
                params={"run_id": run_id},
                json={"contact_info": "jane@example.com"},
            )
            assert first.json()["checks"][0]["status"] == "fail"
            assert second.json()["checks"][0]["status"] == "pass"


async def test_quality_report_reuses_cached_improved_result_for_docx(
    monkeypatch: pytest.MonkeyPatch,
    app_client: tuple[object, CvStore],
) -> None:
    """The two-column .docx's underlying reworded text is expensive (LLM); it must
    reuse the same `improved_result` cache `/cv/{id}/improve` already populates."""
    app, _store = app_client
    monkeypatch.setattr(
        "jobmarket.api.cv_routes.extract_cv_profile_with_fallback", _fake_extract_deterministic
    )
    monkeypatch.setattr(
        "jobmarket.api.quality_routes.generate_cv_quality_report",
        lambda profile, skill_gap, matched_jobs, **kw: _fake_report(contact_ok=True),
    )
    call_count = 0

    def _counting_generate(profile: object, skill_gap: object, matched_jobs: object):
        nonlocal call_count
        call_count += 1
        return _fake_improved_result()

    monkeypatch.setattr("jobmarket.api.quality_routes.generate_improved_cv", _counting_generate)

    async with _fixture() as fixture, async_session_factory() as session:
        run_id = await _add_run(session, fixture)
        await _add_job(session, fixture, "Python Developer", ["Python", "AWS"], run_id)
        await session.commit()

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            upload_resp = await client.post(
                "/cv/upload",
                files={"file": ("cv.txt", _strong_cv_text().encode(), "text/plain")},
            )
            cv_id = upload_resp.json()["cv_id"]

            first = await client.post(f"/cv/{cv_id}/quality-report", params={"run_id": run_id})
            second = await client.post(f"/cv/{cv_id}/quality-report", params={"run_id": run_id})
            assert first.status_code == 200
            assert second.status_code == 200
            assert call_count == 1  # LLM only called once despite two requests


async def test_quality_report_unknown_id_returns_404(app_client: tuple[object, CvStore]) -> None:
    app, _store = app_client
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post("/cv/does-not-exist/quality-report")
    assert resp.status_code == 404


async def test_quality_report_out_of_scope_short_circuits(
    monkeypatch: pytest.MonkeyPatch,
    app_client: tuple[object, CvStore],
) -> None:
    app, _store = app_client
    monkeypatch.setattr(
        "jobmarket.api.cv_routes.extract_cv_profile_with_fallback", _fake_extract_deterministic
    )
    sales_cv_text = (
        "Profile\nSales assistant with strong customer relationships.\n\n"
        "Experience\nJan 2023 - Dec 2023 Sales Assistant - ExampleCo\n"
        "Vente de produits, relation client, service client, encaissement.\n"
    )

    async with _fixture() as fixture, async_session_factory() as session:
        run_id = await _add_run(session, fixture)
        await _add_job(session, fixture, "Python Developer", ["Python"], run_id)
        await session.commit()

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            upload_resp = await client.post(
                "/cv/upload",
                files={"file": ("sales.txt", sales_cv_text.encode(), "text/plain")},
            )
            cv_id = upload_resp.json()["cv_id"]

            resp = await client.post(f"/cv/{cv_id}/quality-report", params={"run_id": run_id})
            assert resp.status_code == 200
            payload = resp.json()
            assert payload["out_of_scope"] is True
            assert payload["docx_base64"] is None
            assert "outside the IT/tech scope" in payload["message"]


async def test_quality_report_improve_failure_surfaces_502(
    monkeypatch: pytest.MonkeyPatch,
    app_client: tuple[object, CvStore],
) -> None:
    """No deterministic fallback exists for the downloadable CV — a total LLM
    failure must be surfaced honestly as an error, never a fabricated document."""
    app, _store = app_client
    monkeypatch.setattr(
        "jobmarket.api.cv_routes.extract_cv_profile_with_fallback", _fake_extract_deterministic
    )
    monkeypatch.setattr(
        "jobmarket.api.quality_routes.generate_cv_quality_report",
        lambda profile, skill_gap, matched_jobs, **kw: _fake_report(contact_ok=True),
    )

    def _raise(profile: object, skill_gap: object, matched_jobs: object) -> ImprovedCvResult:
        raise CvImproverError("LLM call failed: network down")

    monkeypatch.setattr("jobmarket.api.quality_routes.generate_improved_cv", _raise)

    async with _fixture() as fixture, async_session_factory() as session:
        run_id = await _add_run(session, fixture)
        await _add_job(session, fixture, "Python Developer", ["Python", "AWS"], run_id)
        await session.commit()

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            upload_resp = await client.post(
                "/cv/upload",
                files={"file": ("cv.txt", _strong_cv_text().encode(), "text/plain")},
            )
            cv_id = upload_resp.json()["cv_id"]

            resp = await client.post(f"/cv/{cv_id}/quality-report", params={"run_id": run_id})
            assert resp.status_code == 502
            assert "network down" in resp.json()["detail"]
