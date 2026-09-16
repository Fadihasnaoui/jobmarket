"""Integration tests for POST /cv/{id}/improve.

Same DB-backed fixture convention as `test_api_cv.py`. All LLM calls mocked at
`jobmarket.api.improve_routes.generate_improved_cv` — zero network access.
"""

from __future__ import annotations

import base64

import pytest
from httpx import ASGITransport, AsyncClient
from test_api_cv import FakeEncoder, _add_job, _add_run, _fixture, _strong_cv_text

from jobmarket.api.app import create_app
from jobmarket.api.dependencies import get_cv_store, get_upload_encoder
from jobmarket.api.store import CvStore
from jobmarket.cv.extraction import extract_cv_profile
from jobmarket.cv.improver import (
    CvImproverError,
    ImprovedCvResult,
    ImprovedExperienceEntry,
    RecommendedSkill,
)
from jobmarket.cv.llm_extraction import CvExtractionResult
from jobmarket.db.session import async_session_factory


def _fake_extract_deterministic(
    document: object,
    *,
    ontology: object = None,
    skill_ids: object = None,
    client: object = None,
):
    profile = extract_cv_profile(document, ontology=ontology, skill_ids=skill_ids)
    return CvExtractionResult(profile, [], "deterministic_fallback", "forced for test")


def _fake_improved_result(**overrides: object) -> ImprovedCvResult:
    defaults: dict[str, object] = {
        "improved_summary": "Reworded summary using Python and AWS.",
        "improved_experiences": [
            ImprovedExperienceEntry(
                title="Python Developer",
                employer="ExampleCo",
                start_date="Jan 2023",
                end_date="Dec 2023",
                entry_type="job",
                original_evidence="Built APIs with Python and deployed on AWS.",
                improved_description="Engineered APIs in Python and deployed them on AWS.",
            )
        ],
        "skills_section": ["Python", "AWS"],
        "recommended_skills_to_develop": [
            RecommendedSkill(
                canonical_skill="Kubernetes",
                category="tool",
                demand_pct_of_matches=50.0,
                overall_job_demand=10,
            ),
        ],
        "changes_explanation": ["Reworded 1 of 1 experience bullet(s)."],
        "source_degraded": False,
        "source_degraded_reason": None,
    }
    defaults.update(overrides)
    return ImprovedCvResult(**defaults)


@pytest.fixture
def app_client() -> tuple[object, CvStore]:
    app = create_app()
    store = CvStore()
    app.dependency_overrides[get_cv_store] = lambda: store
    app.dependency_overrides[get_upload_encoder] = lambda: FakeEncoder()
    return app, store


async def test_improve_cv_returns_reworded_result_and_docx(
    monkeypatch: pytest.MonkeyPatch,
    app_client: tuple[object, CvStore],
) -> None:
    app, _store = app_client
    monkeypatch.setattr(
        "jobmarket.api.cv_routes.extract_cv_profile_with_fallback", _fake_extract_deterministic
    )
    monkeypatch.setattr(
        "jobmarket.api.improve_routes.generate_improved_cv",
        lambda profile, skill_gap, matched_jobs: _fake_improved_result(),
    )

    async with _fixture() as fixture, async_session_factory() as session:
        run_id = await _add_run(session, fixture)
        await _add_job(session, fixture, "Python Developer", ["Python", "AWS"], run_id)
        await session.commit()

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            upload_resp = await client.post(
                "/cv/upload",
                files={"file": ("cv.txt", _strong_cv_text().encode(), "text/plain")},
            )
            cv_id = upload_resp.json()["cv_id"]

            improve_resp = await client.post(
                f"/cv/{cv_id}/improve",
                params={"run_id": run_id},
                json={"candidate_name": "Jane Doe", "contact_info": "jane@example.com"},
            )
            assert improve_resp.status_code == 200
            payload = improve_resp.json()

            assert payload["out_of_scope"] is False
            assert payload["improved_summary"] == "Reworded summary using Python and AWS."
            assert payload["improved_experiences"][0]["employer"] == "ExampleCo"
            assert payload["improved_experiences"][0]["title"] == "Python Developer"
            assert payload["skills_section"] == ["Python", "AWS"]
            assert payload["recommended_skills_to_develop"][0]["canonical_skill"] == "Kubernetes"
            assert payload["source_degraded"] is False

            # The .docx is real, valid bytes with the given name embedded.
            docx_bytes = base64.b64decode(payload["docx_base64"])
            assert docx_bytes[:2] == b"PK"  # docx is a zip container

            import io

            from docx import Document

            document = Document(io.BytesIO(docx_bytes))
            all_text = "\n".join(p.text for p in document.paragraphs)
            assert "Jane Doe" in all_text
            assert "jane@example.com" in all_text


async def test_improve_cv_caches_llm_result_across_calls(
    monkeypatch: pytest.MonkeyPatch,
    app_client: tuple[object, CvStore],
) -> None:
    """A second call (e.g. regenerating the .docx with a different name) must reuse
    the cached ImprovedCvResult, not re-trigger the LLM."""
    app, _store = app_client
    monkeypatch.setattr(
        "jobmarket.api.cv_routes.extract_cv_profile_with_fallback", _fake_extract_deterministic
    )
    call_count = 0

    def _counting_generate(
        profile: object, skill_gap: object, matched_jobs: object
    ) -> ImprovedCvResult:
        nonlocal call_count
        call_count += 1
        return _fake_improved_result()

    monkeypatch.setattr("jobmarket.api.improve_routes.generate_improved_cv", _counting_generate)

    async with _fixture() as fixture, async_session_factory() as session:
        run_id = await _add_run(session, fixture)
        await _add_job(session, fixture, "Python Developer", ["Python", "AWS"], run_id)
        await session.commit()

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            upload_resp = await client.post(
                "/cv/upload",
                files={"file": ("cv.txt", _strong_cv_text().encode(), "text/plain")},
            )
            cv_id = upload_resp.json()["cv_id"]

            first = await client.post(
                f"/cv/{cv_id}/improve",
                params={"run_id": run_id},
                json={"candidate_name": "Jane Doe"},
            )
            second = await client.post(
                f"/cv/{cv_id}/improve",
                params={"run_id": run_id},
                json={"candidate_name": "John Smith"},
            )
            assert first.status_code == 200
            assert second.status_code == 200
            assert call_count == 1  # LLM only called once despite two requests

            import base64 as b64
            import io

            from docx import Document

            second_docx = Document(io.BytesIO(b64.b64decode(second.json()["docx_base64"])))
            second_text = "\n".join(p.text for p in second_docx.paragraphs)
            assert "John Smith" in second_text
            assert "Jane Doe" not in second_text


async def test_improve_cv_unknown_id_returns_404(app_client: tuple[object, CvStore]) -> None:
    app, _store = app_client
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post("/cv/does-not-exist/improve")
    assert resp.status_code == 404


async def test_improve_cv_out_of_scope_short_circuits(
    monkeypatch: pytest.MonkeyPatch,
    app_client: tuple[object, CvStore],
) -> None:
    app, _store = app_client
    monkeypatch.setattr(
        "jobmarket.api.cv_routes.extract_cv_profile_with_fallback", _fake_extract_deterministic
    )
    sales_cv_text = (
        "Profile\nSales assistant with strong customer relationships.\n\n"
        "Experience\nJan 2023 - Dec 2023 Sales Assistant - ExampleCo\n"
        "Vente de produits, relation client, service client, encaissement.\n"
    )

    async with _fixture() as fixture, async_session_factory() as session:
        run_id = await _add_run(session, fixture)
        await _add_job(session, fixture, "Python Developer", ["Python"], run_id)
        await session.commit()

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            upload_resp = await client.post(
                "/cv/upload",
                files={"file": ("sales.txt", sales_cv_text.encode(), "text/plain")},
            )
            cv_id = upload_resp.json()["cv_id"]

            improve_resp = await client.post(f"/cv/{cv_id}/improve", params={"run_id": run_id})
            assert improve_resp.status_code == 200
            payload = improve_resp.json()
            assert payload["out_of_scope"] is True
            assert payload["docx_base64"] is None
            assert "outside the IT/tech scope" in payload["message"]


async def test_improve_cv_llm_failure_surfaces_502(
    monkeypatch: pytest.MonkeyPatch,
    app_client: tuple[object, CvStore],
) -> None:
    """No deterministic fallback exists for improvement — a total LLM failure must be
    surfaced honestly as an error, never a fabricated 'improved' result."""
    app, _store = app_client
    monkeypatch.setattr(
        "jobmarket.api.cv_routes.extract_cv_profile_with_fallback", _fake_extract_deterministic
    )

    def _raise(profile: object, skill_gap: object, matched_jobs: object) -> ImprovedCvResult:
        raise CvImproverError("LLM call failed: network down")

    monkeypatch.setattr("jobmarket.api.improve_routes.generate_improved_cv", _raise)

    async with _fixture() as fixture, async_session_factory() as session:
        run_id = await _add_run(session, fixture)
        await _add_job(session, fixture, "Python Developer", ["Python", "AWS"], run_id)
        await session.commit()

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            upload_resp = await client.post(
                "/cv/upload",
                files={"file": ("cv.txt", _strong_cv_text().encode(), "text/plain")},
            )
            cv_id = upload_resp.json()["cv_id"]

            improve_resp = await client.post(f"/cv/{cv_id}/improve", params={"run_id": run_id})
            assert improve_resp.status_code == 502
            assert "network down" in improve_resp.json()["detail"]
