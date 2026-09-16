"""Tests for the .docx renderer. Pure rendering logic — no LLM, no network."""

from __future__ import annotations

import io

from docx import Document
from docx.document import Document as DocxDocument

from jobmarket.cv.docx_export import render_improved_cv_docx, render_two_column_cv_docx
from jobmarket.cv.improver import ImprovedCvResult, ImprovedExperienceEntry, RecommendedSkill
from jobmarket.cv.profile import (
    CandidateAttributes,
    CvProfile,
    DocumentPage,
    EducationDetail,
    ExtractionQuality,
    ExtractionVersions,
    ParsedDocument,
)


def _document() -> ParsedDocument:
    text = "x"
    return ParsedDocument(
        filename="cv.txt",
        mime_type="text/plain",
        text=text,
        pages=[DocumentPage(page=None, text=text, start_offset=0, end_offset=len(text))],
        warnings=[],
        file_hash="a" * 64,
        file_size=len(text.encode()),
    )


def _profile(
    *, education: EducationDetail | None = None, confidence_label: str = "high"
) -> CvProfile:
    return CvProfile(
        document=_document(),
        versions=ExtractionVersions(
            parser_version="v1", matcher_version="v1", ontology_version="v1"
        ),
        skills=[],
        candidate=CandidateAttributes(),
        education_detail=education,
        extraction_quality=ExtractionQuality(extraction_confidence_label=confidence_label),
    )


def _result(**overrides: object) -> ImprovedCvResult:
    defaults: dict[str, object] = {
        "improved_summary": "Skilled professional in Python and SQL.",
        "improved_experiences": [
            ImprovedExperienceEntry(
                title="Data Scientist",
                employer="Acme Corp",
                start_date="2020",
                end_date="2023",
                entry_type="job",
                original_evidence="Built pipelines. Improved accuracy by 10%.",
                improved_description="Built pipelines. Improved accuracy by 10%.",
            )
        ],
        "skills_section": ["Python", "SQL"],
        "recommended_skills_to_develop": [
            RecommendedSkill(
                canonical_skill="Docker",
                category="tool",
                demand_pct_of_matches=48.0,
                overall_job_demand=4800,
            ),
        ],
        "changes_explanation": [],
        "source_degraded": False,
        "source_degraded_reason": None,
    }
    defaults.update(overrides)
    return ImprovedCvResult(**defaults)


def _read_back(docx_bytes: bytes) -> DocxDocument:
    return Document(io.BytesIO(docx_bytes))


def _all_text(document: DocxDocument) -> str:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def test_render_produces_valid_docx_with_expected_sections() -> None:
    result = _result()
    profile = _profile(
        education=EducationDetail(education_level="master/mba", institution="MIT")
    )

    docx_bytes = render_improved_cv_docx(result, profile, candidate_name="Jane Doe")
    document = _read_back(docx_bytes)
    text = _all_text(document)

    assert "Jane Doe" in text
    assert "Professional Summary" in text
    assert "Skills" in text
    assert "Python, SQL" in text
    assert "Experience" in text
    assert "Data Scientist" in text
    assert "Acme Corp" in text
    assert "Education" in text
    assert "MIT" in text
    assert "Recommended Skills to Develop for Your Target Roles" in text
    assert "Docker" in text
    assert "48%" in text


def test_missing_name_shows_unmistakable_placeholder() -> None:
    docx_bytes = render_improved_cv_docx(
        _result(), _profile(), candidate_name=None, contact_info=None
    )
    document = _read_back(docx_bytes)

    title_paragraph = document.paragraphs[0]
    assert title_paragraph.text == "[Your Name]"
    assert all(run.italic for run in title_paragraph.runs)


def test_provided_name_and_contact_are_not_italicized_placeholders() -> None:
    docx_bytes = render_improved_cv_docx(
        _result(), _profile(), candidate_name="Jane Doe", contact_info="jane@example.com"
    )
    document = _read_back(docx_bytes)

    title_paragraph = document.paragraphs[0]
    assert title_paragraph.text == "Jane Doe"
    assert not any(run.italic for run in title_paragraph.runs)
    assert "jane@example.com" in _all_text(document)


def test_no_contact_line_when_omitted() -> None:
    docx_bytes = render_improved_cv_docx(
        _result(), _profile(), candidate_name="Jane Doe", contact_info=None
    )
    document = _read_back(docx_bytes)

    # Second paragraph should be the "Note"/heading, not a fabricated contact line.
    assert "@" not in document.paragraphs[1].text


def test_degraded_notice_appears_when_flagged() -> None:
    result = _result(source_degraded=True, source_degraded_reason="Only 1 grounded skill found.")
    docx_bytes = render_improved_cv_docx(result, _profile())
    text = _all_text(_read_back(docx_bytes))

    assert "Note:" in text
    assert "low confidence" in text
    assert "Only 1 grounded skill found." in text


def test_no_degraded_notice_when_not_flagged() -> None:
    result = _result(source_degraded=False)
    docx_bytes = render_improved_cv_docx(result, _profile())
    text = _all_text(_read_back(docx_bytes))

    assert "Note:" not in text


def test_recommended_skills_section_omitted_when_empty() -> None:
    result = _result(recommended_skills_to_develop=[])
    docx_bytes = render_improved_cv_docx(result, _profile())
    document = _read_back(docx_bytes)

    assert "Recommended Skills" not in _all_text(document)
    assert len(document.tables) == 0


def test_experience_description_splits_into_bullet_sentences() -> None:
    result = _result()
    docx_bytes = render_improved_cv_docx(result, _profile())
    document = _read_back(docx_bytes)

    bullet_paragraphs = [
        p for p in document.paragraphs if p.style is not None and p.style.name == "List Bullet"
    ]
    bullet_texts = [p.text for p in bullet_paragraphs]
    assert "Built pipelines." in bullet_texts
    assert "Improved accuracy by 10%." in bullet_texts


def test_no_experience_entries_shows_placeholder() -> None:
    result = _result(improved_experiences=[])
    docx_bytes = render_improved_cv_docx(result, _profile())
    text = _all_text(_read_back(docx_bytes))

    assert "(no experience entries extracted)" in text


def test_no_education_shows_placeholder() -> None:
    docx_bytes = render_improved_cv_docx(_result(), _profile(education=None))
    text = _all_text(_read_back(docx_bytes))

    assert "(no education detected)" in text


# ---------------------------------------------------------------------------
# Two-column template
# ---------------------------------------------------------------------------


def test_two_column_layout_uses_a_single_two_cell_table() -> None:
    docx_bytes = render_two_column_cv_docx(_result(), _profile(), candidate_name="Jane Doe")
    document = _read_back(docx_bytes)

    # The layout table is always added first; a second (1x1) table follows only
    # for the separately-boxed "Recommended Skills" section, when present.
    layout_table = document.tables[0]
    assert len(layout_table.rows) == 1
    assert len(layout_table.rows[0].cells) == 2


def test_two_column_renders_all_expected_sections() -> None:
    result = _result()
    profile = _profile(
        education=EducationDetail(education_level="master/mba", institution="MIT")
    )

    docx_bytes = render_two_column_cv_docx(
        result,
        profile,
        candidate_name="Jane Doe",
        candidate_title="Data Scientist",
        contact_info="jane@example.com | +1 555 0100",
    )
    document = _read_back(docx_bytes)
    text = _all_text(document)

    assert "Jane Doe" in text
    assert "Data Scientist" in text
    assert "CONTACT" in text
    assert "jane@example.com" in text
    assert "+1 555 0100" in text
    assert "SKILLS" in text
    assert "Python" in text
    assert "SQL" in text
    assert "Experience" in text
    assert "Acme Corp" in text
    assert "Education" in text
    assert "MIT" in text
    assert "Recommended Skills to Develop for Your Target Roles" in text
    assert "Docker" in text
    # Certifications has no grounded data source anywhere in the pipeline (see
    # docx_export.py module docstring) — it must never appear, not even as a
    # heading with nothing under it.
    assert "Certification" not in text


def test_two_column_missing_name_shows_unmistakable_placeholder() -> None:
    docx_bytes = render_two_column_cv_docx(_result(), _profile(), candidate_name=None)
    document = _read_back(docx_bytes)

    title_paragraph = document.paragraphs[0]
    assert title_paragraph.text == "[Your Name]"
    assert all(run.italic for run in title_paragraph.runs)


def test_two_column_missing_contact_flags_instead_of_fabricating() -> None:
    docx_bytes = render_two_column_cv_docx(_result(), _profile(), contact_info=None)
    text = _all_text(_read_back(docx_bytes))

    assert "Add your contact info" in text
    assert "@" not in text.split("Add your contact info", 1)[1].split("\n", 1)[0]


def test_two_column_no_title_omits_title_line() -> None:
    result = _result(recommended_skills_to_develop=[])
    docx_bytes = render_two_column_cv_docx(
        result, _profile(), candidate_name="Jane Doe", candidate_title=None
    )
    document = _read_back(docx_bytes)

    # Only the name heading — no title paragraph, no degraded notice, no fabricated
    # subtitle — precedes the two-column table content.
    assert len(document.paragraphs) == 1
    assert document.paragraphs[0].text == "Jane Doe"


def test_two_column_no_skills_shows_placeholder() -> None:
    docx_bytes = render_two_column_cv_docx(_result(skills_section=[]), _profile())
    text = _all_text(_read_back(docx_bytes))

    assert "(none extracted)" in text


def test_two_column_projects_section_omitted_when_no_project_entries() -> None:
    docx_bytes = render_two_column_cv_docx(_result(), _profile())
    text = _all_text(_read_back(docx_bytes))

    assert "Projects" not in text


def test_two_column_projects_section_appears_for_project_entry_type() -> None:
    result = _result(
        improved_experiences=[
            ImprovedExperienceEntry(
                title="Job Recommender",
                employer=None,
                start_date=None,
                end_date=None,
                entry_type="project",
                original_evidence="Built a job recommender using embeddings.",
                improved_description="Built a job recommender using embeddings.",
            )
        ]
    )
    docx_bytes = render_two_column_cv_docx(result, _profile())
    text = _all_text(_read_back(docx_bytes))

    assert "Projects" in text
    assert "Job Recommender" in text
    # A project entry must never bleed into the Experience section.
    experience_index = text.index("Experience")
    projects_index = text.index("Projects")
    assert experience_index < projects_index


def test_two_column_degraded_notice_appears_when_flagged() -> None:
    result = _result(source_degraded=True, source_degraded_reason="Only 1 grounded skill found.")
    docx_bytes = render_two_column_cv_docx(result, _profile())
    text = _all_text(_read_back(docx_bytes))

    assert "Note:" in text
    assert "Only 1 grounded skill found." in text


def test_two_column_recommended_skills_section_omitted_when_empty() -> None:
    result = _result(recommended_skills_to_develop=[])
    docx_bytes = render_two_column_cv_docx(result, _profile())
    document = _read_back(docx_bytes)

    assert "Recommended Skills" not in _all_text(document)
    # Only the two-column layout table remains — no recommended-skills box.
    assert len(document.tables) == 1
