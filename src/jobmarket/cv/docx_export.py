"""Render an `ImprovedCvResult` into a clean, ATS-friendly .docx.

Pure presentation layer: every fact rendered here already passed through
`cv/improver.py`'s grounding guard — this module adds no facts and runs no LLM, it
only lays out already-validated text. It never writes to disk; callers get raw
bytes and decide what to do with them (an API layer streams them straight to the
client, matching the "no raw CV persistence" posture the rest of this app keeps).

The one thing here that is NOT grounded in the CV at all is name/contact: those are
optional, caller-supplied-at-generation-time strings (the profile itself carries no
PII by design — see `cv/profile.py`) — never stored, inserted into the document only,
and rendered as an unmistakable placeholder if left blank rather than a fabricated-
looking value.

`render_two_column_cv_docx` is a second, deliberately singular layout (one polished
template, not a gallery of designs) built from the exact same `ImprovedCvResult` and
the exact same optional-field contract; it adds `candidate_title` alongside the
existing `candidate_name`/`contact_info` (same never-stored posture). It cannot
render a "Certifications" section: nothing upstream extracts certifications as
structured data at all — `cv/llm_extraction.py` explicitly instructs the model to
exclude them from experience — so there is no grounded source to draw from, and
inventing one here would violate the no-fabrication rule. "Projects" is filtered by
`entry_type == "project"`; today's extractors never produce that value, so this
section is honest but currently always empty — a real, not simulated, gap, not a
placeholder.
"""

from __future__ import annotations

import io
import re
from collections.abc import Sequence
from typing import cast

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Length, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from jobmarket.cv.improver import ImprovedCvResult, ImprovedExperienceEntry, RecommendedSkill
from jobmarket.cv.profile import CvProfile

_HEADER_NAME_PLACEHOLDER = "[Your Name]"
_BRAND_COLOR = RGBColor(0x4F, 0x46, 0xE5)  # indigo-600 - matches the frontend brand
_MUTED_COLOR = RGBColor(0x64, 0x74, 0x8B)  # slate-500
_WARNING_COLOR = RGBColor(0x92, 0x40, 0x0E)  # amber-800-ish
_RECOMMENDED_SKILLS_FILL = "FEF3E2"  # light amber - visually distinct box


def render_improved_cv_docx(
    result: ImprovedCvResult,
    profile: CvProfile,
    *,
    candidate_name: str | None = None,
    contact_info: str | None = None,
) -> bytes:
    """Render one .docx and return its raw bytes."""
    document = Document()
    _set_base_styles(document)

    _add_header(document, candidate_name, contact_info)
    if result.source_degraded:
        _add_degraded_notice(document, result.source_degraded_reason)

    _add_heading(document, "Professional Summary")
    document.add_paragraph(result.improved_summary or "(no summary generated)")

    _add_heading(document, "Skills")
    document.add_paragraph(", ".join(result.skills_section) or "(none extracted)")

    _add_heading(document, "Experience")
    if not result.improved_experiences:
        document.add_paragraph("(no experience entries extracted)")
    for entry in result.improved_experiences:
        _add_experience_entry(document, entry)

    _add_heading(document, "Education")
    _add_education(document, profile)

    if result.recommended_skills_to_develop:
        _add_recommended_skills_section(document, result.recommended_skills_to_develop)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _set_base_styles(document: DocxDocument) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for level in (1, 2):
        heading_style = document.styles[f"Heading {level}"]
        heading_style.font.name = "Calibri"
        heading_style.font.color.rgb = _BRAND_COLOR


def _add_heading(document: DocxDocument, text: str) -> None:
    document.add_heading(text, level=1)


def _add_header(
    document: DocxDocument, candidate_name: str | None, contact_info: str | None
) -> None:
    name_text = (candidate_name or "").strip() or _HEADER_NAME_PLACEHOLDER
    heading = document.add_heading(name_text, level=0)
    if not (candidate_name or "").strip():
        for run in heading.runs:
            run.italic = True
            run.font.color.rgb = _MUTED_COLOR
    contact_text = (contact_info or "").strip()
    if contact_text:
        document.add_paragraph(contact_text)


def _add_degraded_notice(document: DocxDocument, reason: str | None) -> None:
    paragraph = document.add_paragraph()
    label = paragraph.add_run("Note: ")
    label.bold = True
    label.font.color.rgb = _WARNING_COLOR
    detail = "The source CV extraction had low confidence"
    if reason:
        detail += f" ({reason})"
    detail += (
        " — some experience or skill details below may be incomplete. "
        "Review carefully before use."
    )
    body = paragraph.add_run(detail)
    body.italic = True
    body.font.color.rgb = _WARNING_COLOR


_SENTENCE_SPLIT = re.compile(r"(?<=[.!?])\s+")


def _split_sentences(text: str) -> list[str]:
    return [part.strip() for part in _SENTENCE_SPLIT.split(text.strip()) if part.strip()]


def _format_dates(start: str | None, end: str | None) -> str:
    if start and end:
        return f"{start} – {end}"
    return start or end or ""


def _add_experience_entry(document: DocxDocument, entry: ImprovedExperienceEntry) -> None:
    paragraph = document.add_paragraph()
    title_run = paragraph.add_run(entry.title or "Untitled role")
    title_run.bold = True
    if entry.employer:
        paragraph.add_run(f" — {entry.employer}").italic = True
    dates = _format_dates(entry.start_date, entry.end_date)
    if dates:
        date_run = paragraph.add_run(f"  ({dates})")
        date_run.font.size = Pt(9)
        date_run.font.color.rgb = _MUTED_COLOR

    sentences = _split_sentences(entry.improved_description)
    if not sentences:
        document.add_paragraph(entry.improved_description, style="List Bullet")
        return
    for sentence in sentences:
        document.add_paragraph(sentence, style="List Bullet")


def _education_line(profile: CvProfile) -> str | None:
    edu = profile.education_detail
    if edu is None:
        return None
    parts = [part for part in (edu.education_level, edu.education_field, edu.institution) if part]
    line = " · ".join(parts) if parts else "Education detected, no further detail."
    if edu.graduation_year:
        line += f" ({edu.graduation_year})"
    return line


def _add_education(document: DocxDocument, profile: CvProfile) -> None:
    document.add_paragraph(_education_line(profile) or "(no education detected)")


def _shade_cell(cell: _Cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _write_skill_line(paragraph: Paragraph, item: RecommendedSkill) -> None:
    name_run = paragraph.add_run(item.canonical_skill)
    name_run.bold = True
    detail = f" — appears in {item.demand_pct_of_matches:.0f}% of your matched jobs"
    if item.overall_job_demand:
        detail += f" ({item.overall_job_demand:,} jobs corpus-wide)"
    paragraph.add_run(detail)


def _add_recommended_skills_section(document: DocxDocument, items: list[RecommendedSkill]) -> None:
    heading = document.add_heading("Recommended Skills to Develop for Your Target Roles", level=1)
    for run in heading.runs:
        run.font.color.rgb = _WARNING_COLOR

    note = document.add_paragraph()
    note_run = note.add_run(
        "These are NOT skills you currently have — they are gaps identified from "
        "real market demand across jobs matching your profile."
    )
    note_run.italic = True
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = _MUTED_COLOR

    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    _shade_cell(cell, _RECOMMENDED_SKILLS_FILL)

    first_paragraph = cell.paragraphs[0]
    first_paragraph.style = document.styles["List Bullet"]
    _write_skill_line(first_paragraph, items[0])
    for item in items[1:]:
        paragraph = cell.add_paragraph(style="List Bullet")
        _write_skill_line(paragraph, item)


# ---------------------------------------------------------------------------
# Two-column template — one deliberately singular design, not a gallery.
# ---------------------------------------------------------------------------

_SIDEBAR_FILL = "EEF0FB"  # very light indigo tint — visually distinct sidebar panel
_MAIN_COLUMN_WIDTH = Inches(4.1)
_SIDEBAR_COLUMN_WIDTH = Inches(2.3)
_CONTACT_SPLIT = re.compile(r"[|\n;]+")


def render_two_column_cv_docx(
    result: ImprovedCvResult,
    profile: CvProfile,
    *,
    candidate_name: str | None = None,
    candidate_title: str | None = None,
    contact_info: str | None = None,
) -> bytes:
    """Render one polished two-column .docx and return its raw bytes.

    Same grounding posture as `render_improved_cv_docx`: every fact here already
    passed `cv/improver.py`'s guard; this function runs no LLM and adds no facts of
    its own. Contact details are shown once, in the sidebar, rather than duplicated
    in the header — the header carries name and title only.
    """
    document = Document()
    _set_base_styles(document)

    _add_two_column_header(document, candidate_name, candidate_title)
    if result.source_degraded:
        _add_degraded_notice(document, result.source_degraded_reason)

    experiences = [e for e in result.improved_experiences if e.entry_type != "project"]
    projects = [e for e in result.improved_experiences if e.entry_type == "project"]

    table = document.add_table(rows=1, cols=2)
    _set_column_widths(table, (_MAIN_COLUMN_WIDTH, _SIDEBAR_COLUMN_WIDTH))
    main_cell, sidebar_cell = table.rows[0].cells
    _set_cell_margins(main_cell, end=200)
    _set_cell_margins(sidebar_cell, top=120, bottom=120, start=200, end=150)
    _shade_cell(sidebar_cell, _SIDEBAR_FILL)

    _add_main_experience(main_cell, experiences)
    _add_main_projects(main_cell, projects)
    _add_main_education(main_cell, profile)

    _add_sidebar_contact(sidebar_cell, contact_info)
    _add_sidebar_skills(sidebar_cell, result.skills_section)

    if result.recommended_skills_to_develop:
        _add_recommended_skills_section(document, result.recommended_skills_to_develop)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_two_column_header(
    document: DocxDocument, candidate_name: str | None, candidate_title: str | None
) -> None:
    name_text = (candidate_name or "").strip() or _HEADER_NAME_PLACEHOLDER
    heading = document.add_heading(name_text, level=0)
    if not (candidate_name or "").strip():
        for run in heading.runs:
            run.italic = True
            run.font.color.rgb = _MUTED_COLOR
    title_text = (candidate_title or "").strip()
    if title_text:
        paragraph = document.add_paragraph()
        run = paragraph.add_run(title_text)
        run.font.size = Pt(12)
        run.italic = True
        run.font.color.rgb = _MUTED_COLOR


def _set_column_widths(table: Table, widths: Sequence[Length]) -> None:
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width
    for column, width in zip(table.columns, widths, strict=True):
        column.width = width


def _set_cell_margins(
    cell: _Cell, *, top: int = 80, bottom: int = 80, start: int = 100, end: int = 100
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _next_paragraph(cell: _Cell) -> Paragraph:
    """The cell's own empty first paragraph if untouched, else a fresh one.

    Every new table cell starts with exactly one empty paragraph; reusing it for
    the first heading avoids leaving a stray blank line above it.
    """
    first = cell.paragraphs[0]
    if not first.runs and not first.text:
        return cast(Paragraph, first)
    return cast(Paragraph, cell.add_paragraph())


def _style_main_heading(paragraph: Paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = _BRAND_COLOR
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)


def _style_sidebar_heading(paragraph: Paragraph, text: str) -> None:
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = _BRAND_COLOR
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)


def _add_two_column_entry(
    cell: _Cell,
    title: str | None,
    employer: str | None,
    start_date: str | None,
    end_date: str | None,
    description: str,
) -> None:
    paragraph = cell.add_paragraph()
    title_run = paragraph.add_run(title or "Untitled role")
    title_run.bold = True
    title_run.font.size = Pt(10.5)
    if employer:
        paragraph.add_run(f" — {employer}").italic = True
    dates = _format_dates(start_date, end_date)
    if dates:
        date_run = paragraph.add_run(f"  ({dates})")
        date_run.font.size = Pt(8.5)
        date_run.font.color.rgb = _MUTED_COLOR

    sentences = _split_sentences(description)
    if not sentences:
        cell.add_paragraph(description, style="List Bullet")
        return
    for sentence in sentences:
        cell.add_paragraph(sentence, style="List Bullet")


def _add_main_experience(cell: _Cell, experiences: list[ImprovedExperienceEntry]) -> None:
    heading = _next_paragraph(cell)
    _style_main_heading(heading, "Experience")
    if not experiences:
        run = cell.add_paragraph().add_run("(no experience entries extracted)")
        run.italic = True
        return
    for entry in experiences:
        _add_two_column_entry(
            cell, entry.title, entry.employer, entry.start_date, entry.end_date,
            entry.improved_description,
        )


def _add_main_projects(cell: _Cell, projects: list[ImprovedExperienceEntry]) -> None:
    """Renders only if a grounded `entry_type == "project"` entry exists.

    No current extractor ever produces that entry_type (see module docstring) — so
    this section is honest, not simulated: it will stay empty until upstream
    extraction is extended to recognize projects as their own category.
    """
    if not projects:
        return
    heading = cell.add_paragraph()
    _style_main_heading(heading, "Projects")
    for entry in projects:
        _add_two_column_entry(
            cell, entry.title, entry.employer, entry.start_date, entry.end_date,
            entry.improved_description,
        )


def _add_main_education(cell: _Cell, profile: CvProfile) -> None:
    heading = cell.add_paragraph()
    _style_main_heading(heading, "Education")
    cell.add_paragraph(_education_line(profile) or "(no education detected)")


def _add_sidebar_contact(cell: _Cell, contact_info: str | None) -> None:
    heading = _next_paragraph(cell)
    _style_sidebar_heading(heading, "Contact")
    contact_text = (contact_info or "").strip()
    if not contact_text:
        run = cell.add_paragraph().add_run(
            "Add your contact info using the optional field before generating your CV."
        )
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = _MUTED_COLOR
        return
    parts = [part.strip() for part in _CONTACT_SPLIT.split(contact_text) if part.strip()]
    for part in parts or [contact_text]:
        run = cell.add_paragraph().add_run(part)
        run.font.size = Pt(9.5)


def _add_sidebar_skills(cell: _Cell, skills: list[str]) -> None:
    heading = cell.add_paragraph()
    _style_sidebar_heading(heading, "Skills")
    if not skills:
        run = cell.add_paragraph().add_run("(none extracted)")
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = _MUTED_COLOR
        return
    for skill in skills:
        run = cell.add_paragraph(style="List Bullet").add_run(skill)
        run.font.size = Pt(9.5)


__all__ = ["render_improved_cv_docx", "render_two_column_cv_docx"]
